"""Régression · inventaire d'exploitation importé d'un tableur (art. 60-62, 119, 122, 124)."""

import io

from openpyxl import load_workbook

from tests_e2e.helpers import (
    declare,
    detail,
    imported,
    inventory_csv,
    inventory_rows,
    inventory_upload,
    issue_codes,
    line,
)
from tools import seed_pilot

OKOUME = "Aucoumea klaineana"


def test_l_inventaire_s_importe_avant_l_autorisation_d_exploiter(world):
    """L'inventaire d'exploitation précède le PAO : il s'importe sur une AAC sans autorisation délivrée (art. 62)."""
    out = inventory_upload(world, inventory_rows(world, [(OKOUME, 96, "oui")]))
    assert out["can_confirm"] and out["summary"]["errors"] == 0


def test_une_aac_fermee_refuse_l_import(world):
    """Une AAC fermée refuse l'import de son inventaire."""
    rows = inventory_rows(world, [(OKOUME, 96, "oui")], aac="hob0")
    out = inventory_upload(world, rows, aac="hob0", expect=422)
    assert "AAC fermée" in detail(out)


def test_l_aac_d_une_autre_organisation_refuse_l_import(world):
    """Une entreprise ne peut pas importer l'inventaire de l'AAC d'une autre organisation."""
    rows = inventory_rows(world, [(OKOUME, 96, "oui")])
    out = inventory_upload(world, rows, actor="cfn.dg", expect=422)
    assert "autre organisation" in detail(out)


def test_la_reference_d_aac_doit_etre_celle_de_l_aac_choisie(world):
    """Chaque ligne porte la référence de l'AAC au plan d'aménagement ; une autre AAC est une erreur (art. 119)."""
    rows = inventory_rows(world, [(OKOUME, 96, "oui"), (OKOUME, 97, "oui")])
    rows[1]["reference_aac"] = "AAC 2026-1"
    codes = issue_codes(inventory_upload(world, rows))
    assert "aac_mismatch" not in codes[0] and "aac_mismatch" in codes[1]


def test_une_tige_a_abattre_sous_le_dme_est_refusee(world):
    """Une tige à abattre sous le DME de l'UFA est une erreur à l'import (art. 124)."""
    out = inventory_upload(world, inventory_rows(world, [(OKOUME, 40, "oui")]))
    assert "below_dme" in issue_codes(out)[0] and not out["can_confirm"]


def test_un_arbre_a_preserver_sous_le_dme_est_accepte(world):
    """Un arbre à préserver (« à abattre : non ») sous le DME est inventorié sans erreur."""
    out = inventory_upload(world, inventory_rows(world, [(OKOUME, 40, "non")]))
    assert out["rows"][0]["level"] == "ok"


def test_une_essence_reservee_ne_s_inscrit_qu_a_preserver(world):
    """Une essence mise en réserve inscrite « à abattre : oui » est refusée ; « non », elle est acceptée."""
    rows = inventory_rows(
        world, [("Baillonella toxisperma", 120, "oui"), ("Baillonella toxisperma", 120, "non")]
    )
    codes = issue_codes(inventory_upload(world, rows))
    assert "reserved_species" in codes[0] and "reserved_species" not in codes[1]


def test_une_tige_sans_position_est_refusee(world):
    """Chaque tige doit être géolocalisée : une ligne sans latitude ou longitude est une erreur."""
    rows = inventory_rows(world, [(OKOUME, 96, "non")])
    rows[0]["latitude"] = ""
    out = inventory_upload(world, rows)
    assert "missing_position" in issue_codes(out)[0] and not out["can_confirm"]


def test_une_tige_hors_de_l_aac_est_refusee(world):
    """Une tige dont la position tombe hors du contour de l'AAC est une erreur."""
    rows = inventory_rows(world, [(OKOUME, 96, "oui")], aac="hob1")
    rows[0]["reference_aac"] = "AAC 2026-2"
    out = inventory_upload(world, rows)
    assert "outside_aac" in issue_codes(out)[0]


def test_un_numero_d_arbre_deja_utilise_dans_l_aac_est_refuse(world):
    """Un numéro d'arbre répété dans le fichier, ou déjà inscrit dans l'AAC, est un doublon refusé."""
    first = inventory_rows(world, [(OKOUME, 96, "oui")], start=8)
    inventory_upload(world, first, step="import", expect=201)
    rows = inventory_rows(world, [(OKOUME, 98, "oui"), (OKOUME, 99, "oui")], start=9)
    rows[1]["numero_arbre"] = rows[0]["numero_arbre"]
    again = inventory_rows(world, [(OKOUME, 97, "oui")])[0]
    rows.append({**again, "numero_arbre": first[0]["numero_arbre"]})
    codes = issue_codes(inventory_upload(world, rows))
    assert "duplicate_number" not in codes[0]
    assert "duplicate_number" in codes[1] and "duplicate_number" in codes[2]


def test_la_confirmation_est_refusee_tant_qu_une_ligne_est_en_erreur(world):
    """L'import n'est confirmé que sans aucune ligne en erreur ; sinon rien n'est inscrit au registre."""
    rows = inventory_rows(world, [(OKOUME, 96, "oui"), (OKOUME, 30, "oui")])
    out = inventory_upload(world, rows, step="import", expect=422)
    assert "1 ligne(s) en erreur" in detail(out)
    assert imported(world, rows[0]["numero_arbre"]) is None


def test_l_import_confirme_inscrit_chaque_tige_a_l_inventaire_de_l_aac(world):
    """L'import confirmé inscrit chaque tige « inventorié » dans l'AAC, avec son numéro et un code attribué."""
    rows = inventory_rows(world, [(OKOUME, 96, "oui"), (OKOUME, 50, "non")], start=2)
    extra = inventory_rows(world, [("Lophira alata", 80, "non")], start=4)[0]
    rows.append({**extra, "numero_arbre": ""})
    out = inventory_upload(world, rows, step="import", expect=201)
    created = out["import"]["tree_codes"]
    assert len(created) == 3 and all(c.startswith("GA-I") for c in created)
    felled = imported(world, rows[0]["numero_arbre"])
    kept = imported(world, rows[1]["numero_arbre"])
    assert felled["inventory_destination"] == "abattre"
    assert kept["inventory_destination"] == "preserver"
    assert felled["aac_id"] == kept["aac_id"] == world.aac["hob2"]


def test_un_arbre_a_preserver_ne_peut_jamais_etre_declare(world):
    """Un arbre inventorié à préserver ne peut jamais être déclaré à l'abattage (art. 122)."""
    rows = inventory_rows(world, [(OKOUME, 95, "non")], aac="hob1", start=3)
    out = inventory_upload(world, rows, step="import", aac="hob1", expect=201)
    assert any(i["code"] == "aac_open" for i in out["file_issues"])
    code = out["import"]["tree_codes"][0]
    refused = declare(world, [{**line(world, OKOUME, 95, i=3), "tree_code": code}], expect=422)
    assert "kept_tree" in detail(refused)


def test_un_arbre_a_preserver_ne_peut_pas_etre_abattu(world):
    """Un arbre à préserver est refusé à l'abattage, quel que soit son statut (art. 122)."""
    rows = inventory_rows(world, [("Lophira alata", 90, "non")], aac="hob1", start=4)
    out = inventory_upload(world, rows, step="import", aac="hob1", expect=201)
    code = out["import"]["tree_codes"][0]
    refused = world.p.post(
        "hob.dg", "/api/felling/events", {"tree_code": code, "event": "abattu"}, expect=409
    )
    assert "préserver" in detail(refused)


def test_une_tige_importee_a_abattre_est_reprise_par_la_declaration(world):
    """Une tige importée « à abattre : oui » est reprise par la déclaration de l'entreprise sans être un doublon."""
    rows = inventory_rows(world, [(OKOUME, 99, "oui")], aac="hob1", start=5)
    out = inventory_upload(world, rows, step="import", aac="hob1", expect=201)
    code = out["import"]["tree_codes"][0]
    declared = declare(world, [{**line(world, OKOUME, 99, i=5), "tree_code": code}])
    assert declared["lines_checked"][0]["violations"] == []


def test_le_classeur_excel_et_la_virgule_decimale_sont_acceptes(world):
    """Le modèle Excel rempli s'importe : virgule décimale, « Oui » en majuscule, feuille « Inventaire »."""
    xlsx = world.p.call("hob.dg", "GET", "/api/inventory/template.xlsx", expect=200)
    wb = load_workbook(io.BytesIO(xlsx))
    assert wb.sheetnames == ["Inventaire", "Exemple", "Lisez-moi"]
    lon, lat = seed_pilot._point(world, "hob2", 6)
    wb["Inventaire"].append(
        ["AAC 2026-2", world.code(), "Okoumé", str(lat).replace(".", ","), lon, "96,5", "Oui"]
    )
    buf = io.BytesIO()
    wb.save(buf)
    out = world.p.call(
        "hob.dg",
        "POST",
        f"/api/inventory/aacs/{world.aac['hob2']}/preview",
        files={"file": ("inventaire.xlsx", buf.getvalue())},
        expect=200,
    )
    assert out["format"] == "xlsx" and out["rows"][0]["diametre_cm"] == 96.5
    assert out["can_confirm"]


def test_un_fichier_au_dela_de_la_limite_de_lignes_est_refuse(world):
    """Un fichier de plus de 5 000 tiges est refusé : l'AAC se découpe en plusieurs fichiers."""
    row = inventory_rows(world, [(OKOUME, 96, "oui")])[0]
    body = inventory_csv([{**row, "numero_arbre": ""}] * 5001)
    out = world.p.call(
        "hob.dg",
        "POST",
        f"/api/inventory/aacs/{world.aac['hob2']}/preview",
        files={"file": ("gros.csv", body, "text/csv")},
        expect=413,
    )
    assert "5000" in detail(out)


def test_une_essence_cites_a_abattre_est_signalee_sans_bloquer(world):
    """Une essence CITES à abattre est signalée en avertissement ; l'import reste possible."""
    out = inventory_upload(world, inventory_rows(world, [("Afzelia bipindensis", 90, "oui")]))
    assert "cites_listed" in issue_codes(out)[0] and out["can_confirm"]


def test_une_colonne_obligatoire_absente_refuse_le_fichier(world):
    """Un fichier sans une colonne obligatoire (ex. « a_abattre ») est refusé avant toute vérification."""
    rows = inventory_rows(world, [(OKOUME, 96, "oui")])
    head = ("reference_aac", "essence", "latitude", "longitude", "diametre_cm")
    out = world.p.call(
        "hob.dg",
        "POST",
        f"/api/inventory/aacs/{world.aac['hob2']}/preview",
        files={"file": ("incomplet.csv", inventory_csv(rows, head=head), "text/csv")},
        expect=422,
    )
    assert "a_abattre" in detail(out)

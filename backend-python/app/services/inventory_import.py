"""Exploitation inventory imported from a spreadsheet (Slice 31).

A company uploads the inventory of one AAC (``.xlsx`` or ``.csv``, French
Excel: ``;`` separator and decimal comma accepted). Every row is checked with
the registry rules — DME of the UFA, reserved species, CITES, outside the AAC,
missing position, duplicate code — before anything is written; the import is
confirmed only when no row is in error, and enters every stem in the registry
(``inventorie``) with its destination: trees to preserve (art. 122) are kept
there and can never be declared nor felled. Each confirmed import is recorded
(who, when, which file, counts) for the Ministry.
"""

from __future__ import annotations

import csv
import io
import math
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime

from app.core import inventory_model as model
from app.core.reserved_species import reserved_species_for
from app.core.species_data import resolve_species
from app.services import compliance, db, dme

CODE_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._/-]{2,39}$")
#: Codes attributed at import (``GA-I<import>-<n>``): refused in a file.
GENERATED_CODE_RE = re.compile(r"^GA-I\d+-", re.IGNORECASE)
TEXT_MAX = 500
SHEET_NAME = "Inventaire"
_VALUE_ALIASES = {
    "destination": {"a_abattre": "abattre", "exploiter": "abattre", "tige_d_avenir": "avenir"},
    "marquage": {"plaque": "qr", "plaque_qr": "qr", "gps": "gps_photo", "gps_photo": "gps_photo"},
}


class ImportFileError(ValueError):
    """The file cannot be read as an inventory (format, size, columns)."""

    def __init__(self, message: str, status: int = 422):
        super().__init__(message)
        self.status = status


@dataclass
class Row:
    line: int
    raw: dict[str, object]
    values: dict[str, object] = field(default_factory=dict)
    errors: list[dict] = field(default_factory=list)
    warnings: list[dict] = field(default_factory=list)

    def issue(self, level: str, code: str, column: str | None = None) -> None:
        item = {"code": code, "label": model.ISSUE_LABELS.get(code, code), "column": column}
        (self.errors if level == "error" else self.warnings).append(item)

    @property
    def key(self) -> str:
        return str(self.values.get("code_arbre") or f"#ligne{self.line}")


def normalise(text: object) -> str:
    """Header / list value key: lower case, no accent, ``_`` between words."""
    raw = unicodedata.normalize("NFKD", str(text or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "_", raw.lower()).strip("_")


def _header_map(headers: list[object]) -> tuple[dict[int, str], list[str]]:
    lookup: dict[str, str] = {}
    for col in model.COLUMNS:
        lookup[col.key] = col.key
        for alias in col.aliases:
            lookup[normalise(alias)] = col.key
    mapping: dict[int, str] = {}
    ignored: list[str] = []
    for i, header in enumerate(headers):
        name = normalise(header)
        if not name:
            continue
        key = lookup.get(name) or lookup.get(re.sub(r"_(cm|m)$", "", name))
        if key is None or key in mapping.values():
            ignored.append(str(header))
        else:
            mapping[i] = key
    return mapping, ignored


def file_format(file_name: str, content: bytes) -> str:
    name = (file_name or "").lower()
    if content[:4] == b"PK\x03\x04" or name.endswith(".xlsx"):
        return "xlsx"
    if name.endswith((".xls", ".ods")) or content[:4] == b"\xd0\xcf\x11\xe0":
        raise ImportFileError(
            "Format non pris en charge : enregistrez le tableau en .xlsx ou .csv", 415
        )
    return "csv"


def _csv_rows(content: bytes) -> list[list[object]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("cp1252", errors="replace")
    first = next((ln for ln in text.splitlines() if ln.strip()), "")
    delimiter = max((";", ",", "\t"), key=first.count)
    return [list(r) for r in csv.reader(io.StringIO(text), delimiter=delimiter)]


def _xlsx_rows(content: bytes) -> list[list[object]]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ImportFileError("Classeur illisible : enregistrez-le de nouveau en .xlsx") from exc
    try:
        ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
        out: list[list[object]] = []
        for values in ws.iter_rows(values_only=True):
            out.append(list(values))
            if len(out) > model.MAX_ROWS + 50:
                break
        return out
    finally:
        wb.close()


def read_table(file_name: str, content: bytes) -> tuple[str, list[str], list[Row]]:
    """Parse the file into rows keyed by model column. Returns (format, ignored, rows)."""
    if not content:
        raise ImportFileError("Fichier vide")
    if len(content) > model.MAX_FILE_BYTES:
        raise ImportFileError(
            f"Fichier trop volumineux (au plus {model.MAX_FILE_BYTES // (1024 * 1024)} Mo)", 413
        )
    fmt = file_format(file_name, content)
    table = _xlsx_rows(content) if fmt == "xlsx" else _csv_rows(content)
    filled = [(i, r) for i, r in enumerate(table, start=1) if any(_blank(v) is False for v in r)]
    if not filled:
        raise ImportFileError("Aucune ligne d'en-tête trouvée")
    _header_line, headers = filled[0]
    mapping, ignored = _header_map(headers)
    missing = [c.key for c in model.COLUMNS if c.required and c.key not in mapping.values()]
    if missing:
        raise ImportFileError("Colonne(s) obligatoire(s) absente(s) : " + ", ".join(missing))
    rows = [
        Row(line, {key: (r[i] if i < len(r) else None) for i, key in mapping.items()})
        for line, r in filled[1:]
    ]
    if not rows:
        raise ImportFileError("Aucune tige dans le fichier")
    if len(rows) > model.MAX_ROWS:
        raise ImportFileError(
            f"Trop de lignes : {model.MAX_ROWS} tiges au plus par fichier — découpez l'AAC en plusieurs fichiers",
            413,
        )
    return fmt, ignored, rows


def _blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def parse_number(value: object) -> float | None:
    """A number typed in Excel or text, decimal comma or point; None if invalid."""
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        number = float(value)
    else:
        text = re.sub(r"[\s  ]", "", str(value))
        if "," in text and "." in text:
            text = (
                text.replace(".", "").replace(",", ".")
                if text.rfind(",") > text.rfind(".")
                else text.replace(",", "")
            )
        else:
            text = text.replace(",", ".")
        try:
            number = float(text)
        except ValueError:
            return None
    return number if math.isfinite(number) else None


def parse_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%Y-%m-%d", "%d/%m/%y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _text(value: object) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def check_values(row: Row, today: date | None = None) -> None:
    """Formats, lists and required values of one row (no database)."""
    today = today or date.today()
    cols = model.by_key()
    for key, col in cols.items():
        raw = row.raw.get(key)
        if _blank(raw):
            if key in ("latitude", "longitude"):
                if not any(e["code"] == "missing_position" for e in row.errors):
                    row.issue("error", "missing_position", key)
            elif col.required:
                row.issue("error", "required", key)
            continue
        if col.kind in ("nombre", "degrés décimaux"):
            number = parse_number(raw)
            if number is None or (col.kind == "nombre" and number <= 0):
                row.issue("error", "invalid_number", key)
                continue
            row.values[key] = number
        elif col.values:
            value = normalise(_text(raw))
            value = _VALUE_ALIASES.get(key, {}).get(value, value)
            if value not in col.values:
                row.issue("error", "invalid_value", key)
                continue
            row.values[key] = value
        elif key == "date_releve":
            parsed = parse_date(raw)
            if parsed is None:
                row.issue("error", "invalid_date", key)
            elif parsed > today:
                row.issue("error", "future_date", key)
            else:
                row.values[key] = parsed
        else:
            text = _text(raw)
            if len(text) > TEXT_MAX:
                row.issue("error", "too_long", key)
                continue
            row.values[key] = text
    code = row.values.get("code_arbre")
    if code is not None:
        if GENERATED_CODE_RE.match(str(code)):
            row.issue("error", "reserved_code", "code_arbre")
        elif not CODE_RE.match(str(code)):
            row.issue("error", "invalid_code", "code_arbre")
    dbh = row.values.get("dhp_cm")
    if dbh is not None and dbh > 400:
        row.issue("error", "invalid_number", "dhp_cm")
    lat, lon = row.values.get("latitude"), row.values.get("longitude")
    if lat is not None and lon is not None and not (-90 <= lat <= 90 and -180 <= lon <= 180):
        row.issue("error", "invalid_position", "latitude")
        del row.values["latitude"]
        del row.values["longitude"]
    row.values.setdefault("marquage", "qr")
    if row.values["marquage"] == "peinture" and not row.values.get("numero_peint"):
        row.issue("error", "painted_number_required", "numero_peint")
    destination = row.values.get("destination")
    species = row.values.get("essence")
    if species:
        profile = resolve_species(str(species))
        if profile is None:
            row.issue("warning", "unknown_species", "essence")
        else:
            row.values["scientific_name"] = profile.scientific_name
            if destination == "abattre" and profile.cites_appendix:
                row.issue("warning", "cites_listed", "essence")
        if destination == "abattre" and reserved_species_for(str(species)) is not None:
            row.issue("error", "reserved_species", "essence")
    if destination == "abattre" and row.values.get("qualite") == model.UNUSABLE_QUALITY:
        row.issue("warning", "unusable_quality", "qualite")


async def aac_context(aac_id: int) -> dict | None:
    row = await db.fetch_one(
        """
        SELECT b.id, b.label, b.status, b.ufa_id, b.geom IS NOT NULL AS has_geom,
               u.label AS ufa_label, t.id AS title_id, t.reference AS title_reference,
               t.holder_org_id
        FROM annual_cut_blocks b
        JOIN management_units u ON u.id = b.ufa_id
        JOIN forest_titles t ON t.id = u.title_id
        WHERE b.id = $1
        """,
        aac_id,
    )
    return dict(row) if row is not None else None


async def check_rows(aac: dict, rows: list[Row]) -> list[dict]:
    """Registry rules over the whole file; returns the file-level warnings."""
    file_issues: list[dict] = []
    if not aac.get("has_geom"):
        file_issues.append(_file_issue("aac_without_geometry"))
    if aac.get("status") in ("autorise", "ouvert"):
        file_issues.append(_file_issue("aac_open"))
    codes = [str(r.values["code_arbre"]) for r in rows if r.values.get("code_arbre")]
    taken = (
        {
            r["tree_code"]
            for r in await db.fetch_all(
                "SELECT tree_code FROM tree_registry WHERE tree_code = ANY($1::text[])", codes
            )
        }
        if codes
        else set()
    )
    painted = [str(r.values["numero_peint"]) for r in rows if r.values.get("numero_peint")]
    painted_taken = (
        {
            r["painted_number"]
            for r in await db.fetch_all(
                "SELECT painted_number FROM tree_registry WHERE aac_id = $1"
                " AND painted_number = ANY($2::text[])",
                aac["id"],
                painted,
            )
        }
        if painted
        else set()
    )
    seen_codes: set[str] = set()
    seen_painted: set[str] = set()
    for row in rows:
        code = row.values.get("code_arbre")
        if code is not None:
            if code in taken or code in seen_codes:
                row.issue("error", "duplicate_code", "code_arbre")
            seen_codes.add(str(code))
        number = row.values.get("numero_peint")
        if number is not None:
            if number in painted_taken or number in seen_painted:
                row.issue("error", "duplicate_painted_number", "numero_peint")
            seen_painted.add(str(number))
    located = [r for r in rows if "latitude" in r.values and "longitude" in r.values]
    points = [(r.key, r.values["longitude"], r.values["latitude"]) for r in located]
    outside = await compliance.lines_outside_aac(aac["id"], points)
    spatial = await compliance.lines_spatial_context(points)
    ufa_dme = await dme.map_for_ufa(aac.get("ufa_id"))
    for row in located:
        if row.key in outside:
            row.issue("error", "outside_aac", "latitude")
        context = spatial.get(row.key) or {}
        if context.get("protected_area"):
            row.issue("warning", "in_protected_area", "latitude")
        if context.get("outside_concession"):
            row.issue("warning", "outside_concession", "latitude")
    for row in rows:
        if row.values.get("destination") != "abattre" or row.values.get("dhp_cm") is None:
            continue
        floor, _source = dme.floor_for(str(row.values.get("essence") or ""), ufa_dme)
        if floor is not None and row.values["dhp_cm"] < floor:
            row.issue("error", "below_dme", "dhp_cm")
            row.values["dme_cm"] = floor
    return file_issues


def _file_issue(code: str, detail: str | None = None) -> dict:
    return {"code": code, "label": model.ISSUE_LABELS.get(code, code), "detail": detail}


def preview(
    aac: dict,
    file_name: str,
    fmt: str,
    ignored: list[str],
    rows: list[Row],
    file_issues: list[dict],
) -> dict:
    if ignored:
        file_issues = [*file_issues, _file_issue("ignored_column", ", ".join(ignored))]
    out_rows = []
    for row in rows:
        level = "error" if row.errors else "warning" if row.warnings else "ok"
        out_rows.append(
            {
                "line": row.line,
                "tree_code": row.values.get("code_arbre"),
                "essence": row.values.get("essence"),
                "scientific_name": row.values.get("scientific_name"),
                "dhp_cm": row.values.get("dhp_cm"),
                "dme_cm": row.values.get("dme_cm"),
                "destination": row.values.get("destination"),
                "numero_peint": row.values.get("numero_peint"),
                "level": level,
                "issues": [{**e, "level": "error"} for e in row.errors]
                + [{**w, "level": "warning"} for w in row.warnings],
            }
        )
    errors = sum(1 for r in out_rows if r["level"] == "error")
    return {
        "aac": {
            "id": aac["id"],
            "label": aac["label"],
            "status": aac["status"],
            "ufa_label": aac["ufa_label"],
            "title_reference": aac["title_reference"],
        },
        "file_name": file_name,
        "format": fmt,
        "model_version": model.MODEL_VERSION,
        "summary": {
            "rows": len(out_rows),
            "ok": sum(1 for r in out_rows if r["level"] == "ok"),
            "warnings": sum(1 for r in out_rows if r["level"] == "warning"),
            "errors": errors,
            "to_fell": sum(1 for r in rows if r.values.get("destination") == "abattre"),
            "kept": sum(1 for r in rows if r.values.get("destination") in model.KEPT_DESTINATIONS),
        },
        "file_issues": file_issues,
        "can_confirm": errors == 0,
        "rows": out_rows,
    }


async def analyse(aac: dict, file_name: str, content: bytes) -> tuple[dict, list[Row]]:
    fmt, ignored, rows = read_table(file_name, content)
    for row in rows:
        check_values(row)
    file_issues = await check_rows(aac, rows)
    return preview(aac, file_name, fmt, ignored, rows, file_issues), rows


async def confirm(
    aac: dict, user: dict, company_id: int | None, result: dict, rows: list[Row]
) -> dict:
    """Record the import and enter every stem in the registry (one statement)."""
    summary = result["summary"]
    header = await db.fetch_one(
        """
        INSERT INTO inventory_imports
            (organization_id, aac_id, title_id, imported_by, file_name, file_format,
             model_version, rows_total, stems_created, to_fell_count, kept_count, warnings_count)
        VALUES ($1, $2, $3, $4, $5, $6, $7, $8, $8, $9, $10, $11)
        RETURNING id, created_at
        """,
        user.get("organization_id"),
        aac["id"],
        aac["title_id"],
        user["id"],
        result["file_name"][:255],
        result["format"],
        model.MODEL_VERSION,
        summary["rows"],
        summary["to_fell"],
        summary["kept"],
        summary["warnings"],
    )

    def col(key: str) -> list:
        return [r.values.get(key) for r in rows]

    try:
        created = await db.fetch_all(
            """
            INSERT INTO tree_registry
                (tree_code, marking_type, geom, species_scientific_name, measured_dbh_cm,
                 company_id, status, permit_number, aac_label, aac_id, title_id,
                 custodian_org_id, inventory_destination, inventory_import_id, painted_number,
                 bole_height_m, inventory_layon, inventory_placette, inventory_quality,
                 surveyed_on, inventory_notes)
            SELECT COALESCE(r.code, 'GA-I' || $1::int || '-' || lpad(r.ord::text, 5, '0')),
                   r.marking, ST_SetSRID(ST_MakePoint(r.lon, r.lat), 4326), r.species, r.dbh,
                   $2, 'inventorie', $3, $4, $5, $6, $7, r.dest, $1, r.painted, r.height,
                   r.layon, r.placette, r.quality, r.surveyed, r.notes
            FROM unnest($8::text[], $9::text[], $10::float8[], $11::float8[], $12::text[],
                        $13::float8[], $14::text[], $15::text[], $16::float8[], $17::text[],
                        $18::text[], $19::text[], $20::date[], $21::text[])
                 WITH ORDINALITY AS r(code, marking, lon, lat, species, dbh, dest, painted,
                                      height, layon, placette, quality, surveyed, notes, ord)
            RETURNING tree_code
            """,
            header["id"],
            company_id,
            aac["title_reference"],
            aac["label"],
            aac["id"],
            aac["title_id"],
            user.get("organization_id"),
            col("code_arbre"),
            col("marquage"),
            col("longitude"),
            col("latitude"),
            [r.values.get("scientific_name") or r.values.get("essence") for r in rows],
            col("dhp_cm"),
            col("destination"),
            col("numero_peint"),
            col("hauteur_fut_m"),
            col("layon"),
            col("placette"),
            col("qualite"),
            col("date_releve"),
            col("observations"),
        )
    except Exception:
        await db.execute("DELETE FROM inventory_imports WHERE id = $1", header["id"])
        raise
    return {
        "id": header["id"],
        "created_at": header["created_at"],
        "stems_created": len(created),
        "tree_codes": [r["tree_code"] for r in created],
    }


async def list_imports(org_id: int | None = None) -> list[dict]:
    rows = await db.fetch_all(
        """
        SELECT i.id, i.organization_id, o.name AS organization_name, i.aac_id,
               b.label AS aac_label, i.title_id, t.reference AS title_reference,
               i.imported_by, u.display_name AS imported_by_name, u.email AS imported_by_email,
               i.file_name, i.file_format, i.model_version, i.rows_total, i.stems_created,
               i.to_fell_count, i.kept_count, i.warnings_count, i.created_at
        FROM inventory_imports i
        JOIN organizations o ON o.id = i.organization_id
        LEFT JOIN annual_cut_blocks b ON b.id = i.aac_id
        LEFT JOIN forest_titles t ON t.id = i.title_id
        LEFT JOIN users u ON u.id = i.imported_by
        WHERE ($1::int IS NULL OR i.organization_id = $1)
        ORDER BY i.created_at DESC, i.id DESC
        LIMIT 200
        """,
        org_id,
    )
    return [dict(r) for r in rows]


EXAMPLE_ROWS: tuple[dict[str, object], ...] = (
    {
        "code_arbre": "",
        "essence": "Aucoumea klaineana",
        "dhp_cm": 87.5,
        "hauteur_fut_m": 18.5,
        "latitude": -0.6912,
        "longitude": 13.2145,
        "destination": "abattre",
        "marquage": "qr",
        "numero_peint": "412",
        "layon": "L12",
        "placette": "P12-04",
        "qualite": "1",
        "date_releve": "14/02/2026",
        "observations": "",
    },
    {
        "essence": "Okoumé",
        "dhp_cm": 52,
        "latitude": -0.6915,
        "longitude": 13.2151,
        "destination": "avenir",
        "numero_peint": "413",
        "layon": "L12",
        "placette": "P12-04",
        "qualite": "1",
        "date_releve": "14/02/2026",
        "observations": "Tige d'avenir sous le DME",
    },
    {
        "essence": "Moabi",
        "dhp_cm": 110,
        "latitude": -0.6921,
        "longitude": 13.2160,
        "destination": "reserve",
        "marquage": "peinture",
        "numero_peint": "414",
        "layon": "L13",
        "placette": "P13-01",
        "date_releve": "15/02/2026",
        "observations": "Essence mise en réserve",
    },
    {
        "essence": "Pterocarpus soyauxii",
        "dhp_cm": 74,
        "latitude": -0.6930,
        "longitude": 13.2172,
        "destination": "semencier",
        "numero_peint": "415",
        "layon": "L13",
        "placette": "P13-02",
        "qualite": "2",
        "date_releve": "15/02/2026",
        "observations": "Semencier à conserver",
    },
)


def _fr_number(value: object) -> str:
    return str(value).replace(".", ",") if isinstance(value, float) else str(value)


def template_csv() -> bytes:
    """``;``-separated template (French Excel), decimal comma, UTF-8 with BOM."""
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=";", lineterminator="\r\n")
    writer.writerow([c.key for c in model.COLUMNS])
    for example in EXAMPLE_ROWS:
        writer.writerow([_fr_number(example.get(c.key, "")) for c in model.COLUMNS])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def template_xlsx() -> bytes:
    """Workbook: « Inventaire » (to fill), « Exemple », « Lisez-moi »."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="4B5320")
    keys = [c.key for c in model.COLUMNS]
    wb = Workbook()
    sheet = wb.active
    sheet.title = SHEET_NAME
    example = wb.create_sheet("Exemple")
    for ws in (sheet, example):
        ws.append(keys)
        for cell in ws[1]:
            cell.font, cell.fill = head_font, head_fill
        ws.freeze_panes = "A2"
        for i, col in enumerate(model.COLUMNS, start=1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = max(12, len(col.key) + 4)
    for row in EXAMPLE_ROWS:
        example.append([row.get(k, None) or None for k in keys])
    last = model.MAX_ROWS + 1
    for i, col in enumerate(model.COLUMNS, start=1):
        if not col.values:
            continue
        letter = sheet.cell(1, i).column_letter
        rule = DataValidation(
            type="list", formula1='"' + ",".join(col.values) + '"', allow_blank=True
        )
        rule.error = "Valeur hors liste : " + ", ".join(col.values)
        sheet.add_data_validation(rule)
        rule.add(f"{letter}2:{letter}{last}")
    readme = wb.create_sheet("Lisez-moi")
    readme.append(["Inventaire d'exploitation d'une AAC — modèle d'import (proposition)"])
    readme["A1"].font = Font(bold=True, size=14)
    readme.append([f"Version du modèle : {model.MODEL_VERSION}"])
    readme.append([model.HELP_TEXT])
    readme.append(
        [
            f"Remplir la feuille « {SHEET_NAME} » : une ligne par tige, {model.MAX_ROWS} tiges au plus"
            f" par fichier ({model.MAX_FILE_BYTES // (1024 * 1024)} Mo). La feuille « Exemple » montre"
            " une ligne de chaque destination ; elle n'est jamais importée."
        ]
    )
    readme.append([])
    readme.append(
        ["Colonne", "Obligatoire", "Format", "Exemple", "Explication", "Fondement", "Valeurs"]
    )
    for cell in readme[6]:
        cell.font, cell.fill = head_font, head_fill
    for col in model.COLUMNS:
        readme.append(
            [
                col.key,
                "oui" if col.required else "non",
                col.kind,
                col.example,
                col.help,
                col.legal,
                ", ".join(col.values),
            ]
        )
    for letter, width in zip("ABCDEFG", (18, 12, 16, 20, 60, 50, 32), strict=True):
        readme.column_dimensions[letter].width = width
    for row in readme.iter_rows(min_row=7):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    readme["A3"].alignment = Alignment(wrap_text=True)
    readme["A4"].alignment = Alignment(wrap_text=True)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()

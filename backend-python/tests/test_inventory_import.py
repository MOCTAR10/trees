"""Inventory import from a spreadsheet (Slice 31) — parsing, row checks, templates, access."""

import asyncio
import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.config import get_settings
from app.core import inventory_model as model
from app.core import security
from app.main import app
from app.services import compliance, db, dme, inventory_import

AAC = {
    "id": 7,
    "label": "PILOTE-AAC 2026-2",
    "status": "soumis",
    "ufa_id": 3,
    "has_geom": True,
    "ufa_label": "UFA 1",
    "title_id": 2,
    "title_reference": "CFAD-01",
    "holder_org_id": 10,
}

HEAD = "reference_aac;numero_arbre;essence;latitude;longitude;diametre_cm;a_abattre"


def _csv(*lines: str, head: str = HEAD) -> bytes:
    return ("\n".join([head, *lines]) + "\n").encode("utf-8")


def _line(
    number: str = "", species: str = "Okoumé", dbh: str = "96", answer: str = "oui", **kw
) -> str:
    lat, lon = kw.get("lat", "-0,77"), kw.get("lon", "13,64")
    return ";".join([kw.get("aac", "AAC 2026-2"), number, species, lat, lon, dbh, answer])


@pytest.fixture
def fake_rules(monkeypatch):
    state = {"outside": set(), "numbers": set(), "spatial": {}, "ufa_dme": {}}

    async def outside(_aac_id, points):
        return {p[0] for p in points if p[0] in state["outside"]}

    async def spatial(points):
        return {p[0]: state["spatial"][p[0]] for p in points if p[0] in state["spatial"]}

    async def fetch_all(sql, *args):
        if "painted_number = any" in sql.lower():
            return [{"painted_number": n} for n in args[1] if n in state["numbers"]]
        return []

    async def ufa_map(_ufa_id):
        return state["ufa_dme"]

    monkeypatch.setattr(compliance, "lines_outside_aac", outside)
    monkeypatch.setattr(compliance, "lines_spatial_context", spatial)
    monkeypatch.setattr(dme, "map_for_ufa", ufa_map)
    monkeypatch.setattr(db, "fetch_all", fetch_all)
    return state


def _analyse(content: bytes, name: str = "inv.csv", aac: dict | None = None) -> dict:
    result, _rows = asyncio.run(inventory_import.analyse(aac or AAC, name, content))
    return result


def _codes(result: dict) -> list[list[str]]:
    return [[i["code"] for i in r["issues"]] for r in result["rows"]]


def test_model_has_only_the_critical_columns():
    assert [c.key for c in model.COLUMNS] == [
        "reference_aac",
        "numero_arbre",
        "essence",
        "latitude",
        "longitude",
        "diametre_cm",
        "a_abattre",
        "observations",
    ]
    assert {c.key for c in model.COLUMNS if not c.required} == {"numero_arbre", "observations"}


def test_numbers_accept_decimal_comma_and_spaces():
    assert inventory_import.parse_number("87,5") == 87.5
    assert inventory_import.parse_number("-0,6912") == -0.6912
    assert inventory_import.parse_number("1 234,5") == 1234.5
    assert inventory_import.parse_number("1.234,5") == 1234.5
    assert inventory_import.parse_number("1,234.5") == 1234.5
    assert inventory_import.parse_number(96) == 96.0
    assert inventory_import.parse_number("abc") is None
    assert inventory_import.parse_number("nan") is None


def test_headers_matched_by_name_alias_case_and_accents():
    mapping, ignored = inventory_import._header_map(
        ["Réf AAC", "Numéro", "Espèce", "Lat", "Longitude", "DHP (cm)", "À abattre", "Couleur"]
    )
    assert set(mapping.values()) == {
        "reference_aac",
        "numero_arbre",
        "essence",
        "latitude",
        "longitude",
        "diametre_cm",
        "a_abattre",
    }
    assert ignored == ["Couleur"]


def test_aac_reference_must_name_the_chosen_aac():
    assert inventory_import.same_aac("PILOTE-AAC 2026-2", "PILOTE-AAC 2026-2")
    assert inventory_import.same_aac("aac 2026-2", "PILOTE-AAC 2026-2")
    assert not inventory_import.same_aac("AAC 2026-21", "PILOTE-AAC 2026-2")
    assert not inventory_import.same_aac("AAC 2025-1", "PILOTE-AAC 2026-2")


def test_csv_separator_and_windows_encoding_detected(fake_rules):
    body = (
        "reference_aac,essence,latitude,longitude,diametre_cm,a_abattre\n"
        "AAC 2026-2,Okoumé,-0.77,13.64,96,oui\n"
    )
    result = _analyse(body.encode("cp1252"))
    assert result["rows"][0]["essence"] == "Okoumé" and result["can_confirm"]
    tab = _analyse(_csv(_line()).replace(b";", b"\t"))
    assert tab["rows"][0]["diametre_cm"] == 96


def test_missing_required_column_refuses_the_file(fake_rules):
    with pytest.raises(inventory_import.ImportFileError, match="a_abattre"):
        _analyse(_csv("AAC 2026-2;1;Okoumé;-0,77;13,64;96", head=HEAD.rsplit(";", 1)[0]))


def test_row_limit_and_file_size(fake_rules, monkeypatch):
    monkeypatch.setattr(model, "MAX_ROWS", 2)
    with pytest.raises(inventory_import.ImportFileError) as err:
        _analyse(_csv(*[_line()] * 3))
    assert err.value.status == 413
    monkeypatch.setattr(model, "MAX_FILE_BYTES", 10)
    with pytest.raises(inventory_import.ImportFileError, match="volumineux"):
        _analyse(_csv(_line()))


def test_old_excel_format_is_refused():
    with pytest.raises(inventory_import.ImportFileError) as err:
        inventory_import.read_table("inventaire.xls", b"\xd0\xcf\x11\xe0rest")
    assert err.value.status == 415


def test_row_rules(fake_rules):
    fake_rules["ufa_dme"] = {"Aucoumea klaineana": {"dme_cm": 75, "source": "arrêté"}}
    fake_rules["outside"] = {"#ligne9"}
    fake_rules["numbers"] = {"77"}
    result = _analyse(
        _csv(
            _line("1"),
            _line("2", dbh="60"),
            _line("3", dbh="60", answer="non"),
            _line("4", species="Moabi", dbh="120"),
            _line("5", species="Moabi", dbh="120", answer="non"),
            _line("6", lat=""),
            _line("7", aac="AAC 2025-1"),
            _line("8"),
            _line("77"),
            _line("9", dbh="abc", answer="peut-être"),
            _line("10", species="Arbre inconnu"),
            _line("11", lat="-95"),
            _line("12", answer="O"),
        )
    )
    codes = _codes(result)
    assert codes[0] == []
    assert codes[1] == ["below_dme"] and result["rows"][1]["dme_cm"] == 75
    assert codes[2] == []
    assert "reserved_species" in codes[3]
    assert codes[4] == []
    assert codes[5] == ["missing_position"]
    assert codes[6] == ["aac_mismatch"]
    assert codes[7] == ["outside_aac"]
    assert codes[8] == ["duplicate_number"]
    assert set(codes[9]) == {"invalid_number", "invalid_value"}
    assert codes[10] == ["unknown_species"] and result["rows"][10]["level"] == "warning"
    assert codes[11] == ["invalid_position"]
    assert codes[12] == [] and result["rows"][12]["destination"] == "abattre"
    assert result["rows"][2]["destination"] == "preserver"
    assert result["summary"]["kept"] == 2 and not result["can_confirm"]


def test_tree_number_repeated_in_the_file(fake_rules):
    result = _analyse(_csv(_line("1"), _line("1"), _line(""), _line("")))
    assert _codes(result) == [[], ["duplicate_number"], [], []]


def test_warnings_do_not_block(fake_rules):
    fake_rules["spatial"] = {"#ligne2": {"protected_area": "Parc", "outside_concession": True}}
    result = _analyse(
        _csv(
            _line(),
            _line(species="Afzelia bipindensis", dbh="90"),
            _line(species="Afzelia bipindensis", dbh="90", answer="non"),
        )
    )
    codes = _codes(result)
    assert set(codes[0]) == {"in_protected_area", "outside_concession"}
    assert codes[1] == ["cites_listed"] and codes[2] == []
    assert result["can_confirm"]


def test_file_level_warnings(fake_rules):
    result = _analyse(
        _csv(_line() + ";rouge", head=HEAD + ";couleur"),
        aac={**AAC, "status": "ouvert", "has_geom": False},
    )
    assert {i["code"] for i in result["file_issues"]} == {
        "aac_open",
        "aac_without_geometry",
        "ignored_column",
    }


def test_xlsx_template_has_three_sheets_and_reads_back(fake_rules):
    content = inventory_import.template_xlsx()
    wb = load_workbook(io.BytesIO(content))
    assert wb.sheetnames == ["Inventaire", "Exemple", "Lisez-moi"]
    assert [c.value for c in wb["Inventaire"][1]] == [c.key for c in model.COLUMNS]
    readme = [r[0] for r in wb["Lisez-moi"].iter_rows(values_only=True)]
    assert all(c.key in readme for c in model.COLUMNS)
    with pytest.raises(inventory_import.ImportFileError, match="Aucune tige"):
        _analyse(content, "modele.xlsx")
    wb["Inventaire"].append(["AAC 2026-2", 412, "Okoumé", -0.77, 13.64, 96.5, "Oui"])
    buf = io.BytesIO()
    wb.save(buf)
    result = _analyse(buf.getvalue(), "inventaire.xlsx")
    row = result["rows"][0]
    assert result["format"] == "xlsx" and row["diametre_cm"] == 96.5
    assert row["numero_arbre"] == "412" and result["can_confirm"]


def test_csv_template_uses_semicolon_bom_and_decimal_comma(fake_rules):
    content = inventory_import.template_csv()
    assert content.startswith(b"\xef\xbb\xbfreference_aac;numero_arbre;essence")
    assert b"87,5" in content
    result = _analyse(content)
    assert result["summary"]["rows"] == len(inventory_import.EXAMPLE_ROWS)
    assert result["summary"]["kept"] == 3


ROLE_IDS = {"admin": 1, "company": 2, "operator": 4, "cooperative": 6}


def _token(role: str) -> dict:
    uid = ROLE_IDS[role]
    token = security.create_access_token(f"{role}@x", role, uid, get_settings().jwt_secret, 3600)
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture
def client(monkeypatch):
    async def fetch_one(sql, *args):
        if "from users" in sql.lower():
            return {
                "id": args[0],
                "email": "co@x",
                "role": next(r for r, i in ROLE_IDS.items() if i == args[0]),
                "is_active": True,
                "organization_id": 10,
                "company_id": None,
                "cooperative_id": None,
                "display_name": "Co",
            }
        return None

    monkeypatch.setattr(db, "fetch_one", fetch_one)
    return TestClient(app)


def test_model_and_templates_open_to_company_and_ministry_only(client):
    resp = client.get("/api/inventory/model", headers=_token("company"))
    assert resp.json()["max_rows"] == model.MAX_ROWS
    resp = client.get("/api/inventory/template.xlsx", headers=_token("admin"))
    assert resp.content.startswith(b"PK")
    for role in ("operator", "cooperative"):
        assert client.get("/api/inventory/model", headers=_token(role)).status_code == 403
        assert client.get("/api/inventory/template.csv", headers=_token(role)).status_code == 403


def test_preview_refused_on_closed_aac_but_allowed_without_authorization(
    client, monkeypatch, fake_rules
):
    async def context(_aac_id):
        return AAC

    gates = {"value": "aac_closed"}

    async def gate(_aac_id, _org):
        return gates["value"]

    monkeypatch.setattr(inventory_import, "aac_context", context)
    monkeypatch.setattr(compliance, "authorization_gate", gate)
    files = {"file": ("inv.csv", _csv(_line()), "text/csv")}
    resp = client.post("/api/inventory/aacs/7/preview", files=files, headers=_token("company"))
    assert resp.status_code == 422 and resp.json()["detail"]["code"] == "aac_closed"
    gates["value"] = "no_authorization"
    resp = client.post("/api/inventory/aacs/7/preview", files=files, headers=_token("company"))
    assert resp.status_code == 200 and resp.json()["can_confirm"]
    resp = client.post("/api/inventory/aacs/7/preview", files=files, headers=_token("admin"))
    assert resp.status_code == 403


def test_confirm_refused_while_a_row_is_in_error(client, monkeypatch, fake_rules):
    async def context(_aac_id):
        return AAC

    async def gate(_aac_id, _org):
        return None

    monkeypatch.setattr(inventory_import, "aac_context", context)
    monkeypatch.setattr(compliance, "inventory_gate", gate)
    files = {"file": ("inv.csv", _csv(_line(lat="", lon="")), "text/csv")}
    resp = client.post("/api/inventory/aacs/7/import", files=files, headers=_token("company"))
    assert resp.status_code == 422
    assert resp.json()["detail"]["rows"][0]["issues"][0]["code"] == "missing_position"


def test_inventory_gate_ignores_missing_authorization(monkeypatch):
    async def gate(_aac_id, _org):
        return gates.pop(0)

    gates = ["no_authorization", "startup_not_notified", "aac_closed", "title_not_exploitable"]
    monkeypatch.setattr(compliance, "authorization_gate", gate)
    results = [asyncio.run(compliance.inventory_gate(1, 1)) for _ in range(4)]
    assert results == [None, None, "aac_closed", "title_not_exploitable"]
